import io
from pathlib import Path
from fastapi import UploadFile, HTTPException

class DocumentParser:
    """
    Parses plain text from text payloads or uploaded files (PDF, DOCX, TXT, MD).
    """
    MAX_FILE_SIZE = 15 * 1024 * 1024  # 15 MB limit
    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".text"}

    @staticmethod
    async def parse_file(file: UploadFile) -> str:
        filename = (file.filename or "").strip()
        file_ext = Path(filename).suffix.lower() if filename else ""
        
        if not file_ext or file_ext not in DocumentParser.ALLOWED_EXTENSIONS:
            allowed = ", ".join(sorted(DocumentParser.ALLOWED_EXTENSIONS))
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file format '{file_ext}'. Allowed formats: {allowed}"
            )

        content_bytes = await file.read()
        if len(content_bytes) > DocumentParser.MAX_FILE_SIZE:
            max_mb = DocumentParser.MAX_FILE_SIZE // (1024 * 1024)
            raise HTTPException(
                status_code=413,
                detail=f"File exceeds maximum size limit of {max_mb}MB."
            )
            
        if not content_bytes or len(content_bytes.strip()) == 0:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")

        if filename.endswith(".pdf"):
            return DocumentParser._parse_pdf(content_bytes)
        elif filename.endswith(".docx"):
            return DocumentParser._parse_docx(content_bytes)
        else:
            return DocumentParser._parse_text(content_bytes)

    @staticmethod
    def _parse_pdf(content_bytes: bytes) -> str:
        extracted_pages = []
        
        # Method 1: Try pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            for page in reader.pages:
                text = page.extract_text()
                if text and text.strip():
                    extracted_pages.append(text.strip())
        except Exception as e:
            print(f"pypdf extraction failed: {e}")

        # Method 2: Fallback to PyPDF2 if pypdf yielded nothing
        if not extracted_pages:
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(content_bytes))
                for page in reader.pages:
                    text = page.extract_text()
                    if text and text.strip():
                        extracted_pages.append(text.strip())
            except Exception as e:
                print(f"PyPDF2 extraction failed: {e}")

        # Method 3: Fallback to raw string extraction
        if not extracted_pages:
            raw = DocumentParser._parse_text(content_bytes)
            return raw.encode('utf-8', 'ignore').decode('utf-8', 'ignore')

        full_text = "\n\n".join(extracted_pages)
        return full_text.encode('utf-8', 'ignore').decode('utf-8', 'ignore')


    @staticmethod
    def _parse_docx(content_bytes: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            lines = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    lines.append(paragraph.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            lines.append(cell.text.strip())
            return "\n".join(lines)
        except Exception:
            return DocumentParser._parse_text(content_bytes)

    @staticmethod
    def _parse_text(content_bytes: bytes) -> str:
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                return content_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content_bytes.decode("utf-8", errors="ignore")


